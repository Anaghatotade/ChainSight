"""
Step 9 - export dashboard-ready files for Power BI / Excel / Tableau.

Writes into outputs/dashboard/:
  * CSV star schema: dim_item, dim_store, dim_week, dim_date, fact_store_dept_week, fact_daily_store_dept,
    plus the KPI / forecast / recommendation tables
  * M5_Retail_Dashboard_Data.xlsx - the smaller tables as formatted sheets (one workbook to open in Excel)
The PNG charts in outputs/charts/ can be dropped straight into PowerPoint / Power BI as images.

Power BI tip: Get Data -> Text/CSV -> pick these files; relate the facts to the dims on item_id / store_id / week_idx / date.
The item-level weekly fact (6.6 million rows) is NOT exported: it is in the SQLite database (data/processed/m5_retail.db).

Run:  python src/09_export_dashboard.py
"""
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config as cfg
import panel as pn

PCT_HINTS = ("share", "wape", "bias", "pct", "rate", "_cv")
MONEY_HINTS = ("revenue", "value")


def store_dept_week_fact(panel):
    """units / revenue / on-sale item count for every store x department x week (19,110 rows) - small enough for any BI tool."""
    s = panel.series
    on, rev = panel.onsale, panel.revenue()
    keys = s["store_id"] + "|" + s["dept_id"]
    parts = []
    for key, idx in keys.groupby(keys).groups.items():
        idx = np.asarray(idx)
        store, dept = key.split("|")
        parts.append(pd.DataFrame({"store_id": store, "dept_id": dept, "week_idx": np.arange(panel.n_weeks),
                                   "units": panel.units[idx].sum(axis=0), "revenue": rev[idx].sum(axis=0).round(2),
                                   "items_on_sale": on[idx].sum(axis=0)}))
    return pd.concat(parts, ignore_index=True)


def main():
    cfg.ensure_dirs()
    T = cfg.TABLES
    panel = pn.load_panel()
    item = pd.read_csv(T / "kpi_item_kpis.csv")
    item_state = pd.read_csv(T / "kpi_item_state_kpis.csv")

    # ------------------------------------------------------------ star schema CSVs
    dim_item = item[["item_id", "dept_id", "cat_id", "abc_class", "avg_selling_price", "zero_week_share"]]
    dim_store = panel.series.drop_duplicates("store_id")[["store_id", "state_id"]].sort_values("store_id")
    dim_week = pd.read_csv(cfg.DATA_PROCESSED / "weeks.csv")
    dim_date = pd.read_csv(cfg.DATA_PROCESSED / "dim_date.csv")
    fact = store_dept_week_fact(panel)
    daily = pd.read_csv(cfg.DATA_PROCESSED / "daily_store_dept.csv")

    # test-week forecasts aggregated to category x week (the item-level file has 430k rows)
    pred = pd.read_csv(cfg.DATA_PROCESSED / "forecast_predictions.csv.gz")
    models = ["naive", "ma", "ses", "snaive", "linear", "mlp"]
    fc_cat_week = pred.groupby(["split", "category", "week_idx", "week_start"], as_index=False)[["actual"] + models].sum().round(1)

    csvs = {"dim_item": dim_item, "dim_store": dim_store, "dim_week": dim_week, "dim_date": dim_date,
            "fact_store_dept_week": fact, "fact_daily_store_dept": daily, "forecast_category_week": fc_cat_week}
    for name in ["kpi_overall", "kpi_item_state_kpis", "kpi_abc_summary", "kpi_abc_by_category", "kpi_demand_class_summary",
                 "kpi_policy_assumed_by_abc", "kpi_policy_assumed_by_cv_band", "kpi_price_cut_uplift", "kpi_slow_movers_C_items",
                 "eda_category_summary", "eda_department_summary", "eda_state_summary", "eda_store_summary", "eda_yearly_summary",
                 "eda_weekly_totals", "eda_weekday_index_category", "forecast_model_comparison", "forecast_segment_wape",
                 "forecast_split_summary", "error_analysis_top100", "error_analysis_tally", "error_analysis_tags", "policy_sim_frontier",
                 "recommendations", "cleaning_log", "data_quality_summary"]:
        csvs[name] = pd.read_csv(T / f"{name}.csv")
    csvs["eda_seasonality_index_dept_month"] = pd.read_csv(T / "eda_seasonality_index_dept_month.csv").rename(columns={"Unnamed: 0": "dept_id"})
    for name, df in csvs.items():
        df.to_csv(cfg.DASHBOARD / f"{name}.csv", index=False, date_format="%Y-%m-%d")

    # ------------------------------------------------------------ Excel workbook (smaller tables only)
    overall = csvs["kpi_overall"].iloc[0]
    definitions = {
        "total_revenue": ("Revenue = units sold x weekly sell price (price is the only money field in M5; no cost data)", "#,##0"),
        "total_units": ("Units sold in the 273 full weeks used", "#,##0"),
        "avg_selling_price": ("Revenue / units", "0.00"),
        "n_items": ("Distinct items (products)", "#,##0"),
        "n_stores": ("Stores (3 states)", "#,##0"),
        "n_weeks": ("Full weeks of history", "#,##0"),
        "on_sale_series_weeks": ("Item-store-weeks in which the item had a price (was on sale)", "#,##0"),
        "zero_sales_week_pct": ("Percent of on-sale item-store-weeks with zero units (NOT the same as a stockout - M5 has no stock data)", "0.0")}
    kpi_summary = pd.DataFrame({"kpi": list(definitions), "value": [overall[k] for k in definitions],
                                "definition": [v[0] for v in definitions.values()]})
    sheets = {
        "KPI_Summary": kpi_summary, "ABC_Summary": csvs["kpi_abc_summary"], "ABC_by_Category": csvs["kpi_abc_by_category"],
        "Demand_Classes": csvs["kpi_demand_class_summary"], "By_Category": csvs["eda_category_summary"],
        "By_Department": csvs["eda_department_summary"], "By_State": csvs["eda_state_summary"], "By_Store": csvs["eda_store_summary"],
        "Yearly": csvs["eda_yearly_summary"], "Seasonality_Dept": csvs["eda_seasonality_index_dept_month"],
        "Weekday_Index": csvs["eda_weekday_index_category"], "Price_Cut_Uplift": csvs["kpi_price_cut_uplift"],
        "Safety_Stock_by_ABC": csvs["kpi_policy_assumed_by_abc"], "Safety_Stock_by_CV": csvs["kpi_policy_assumed_by_cv_band"],
        "Item_Profile": item_state, "Slow_Movers_C": csvs["kpi_slow_movers_C_items"],
        "Forecast_Comparison": csvs["forecast_model_comparison"], "Forecast_Segments": csvs["forecast_segment_wape"],
        "Error_Top100": csvs["error_analysis_top100"], "Error_Tally": csvs["error_analysis_tally"],
        "Policy_Simulation": csvs["policy_sim_frontier"], "Recommendations": csvs["recommendations"], "Cleaning_Log": csvs["cleaning_log"]}
    what = {
        "KPI_Summary": "Headline KPIs with definitions", "ABC_Summary": "What A / B / C items contribute (revenue-based)",
        "ABC_by_Category": "ABC mix per category", "Demand_Classes": "Smooth / erratic / intermittent / lumpy demand (Syntetos-Boylan)",
        "By_Category": "Units, revenue, zero-sales share by category", "By_Department": "Same by department", "By_State": "Same by state",
        "By_Store": "Same by store", "Yearly": "Yearly totals (2011 and 2016 are partial years)", "Seasonality_Dept": "Monthly seasonal index by department (1.0 = normal)",
        "Weekday_Index": "Weekday index by category (1.0 = average day)", "Price_Cut_Uplift": "Units in price-cut weeks vs the previous 8 weeks (association, not causal)",
        "Safety_Stock_by_ABC": "ASSUMPTION-BASED safety stock by ABC class (1-week lead time)", "Safety_Stock_by_CV": "ASSUMPTION-BASED safety stock by variability band",
        "Item_Profile": "One row per item x state: demand mean / CV / class / assumed safety stock", "Slow_Movers_C": "C items with many zero-sales weeks",
        "Forecast_Comparison": "Baselines vs linear vs MLP: WAPE, RMSE, bias", "Forecast_Segments": "WAPE by category / state / ABC / price-cut / event week",
        "Error_Top100": "100 largest dev-set errors with rule-based cause tags (fill manual_cause!)", "Error_Tally": "Count of the top errors by primary cause",
        "Policy_Simulation": "ASSUMPTION-BASED replay: fill rate vs inventory by forecast model", "Recommendations": "Business recommendations (generated)",
        "Cleaning_Log": "Every cleaning rule and rows affected"}
    readme = pd.DataFrame({"sheet": ["README"] + list(sheets), "what_it_contains": ["Data: Kaggle M5 Forecasting - Accuracy (Walmart). Generated by src/09_export_dashboard.py"] + [what[k] for k in sheets]})
    path = cfg.DASHBOARD / "M5_Retail_Dashboard_Data.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        readme.to_excel(xw, sheet_name="README", index=False)       # sheet_name as keyword: works in pandas 2 and 3
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
        header_fill = PatternFill("solid", fgColor="1F4E79")
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
                cell.fill = header_fill
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for ci, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), start=1):
                name = str(col[0].value or "").lower()
                width = min(60, max(10, max(len(str(c.value)) if c.value is not None else 0 for c in col[:60]) + 2))
                if name in ("finding", "action", "what_it_contains", "definition", "rule", "detail"):
                    width = 70
                ws.column_dimensions[get_column_letter(ci)].width = width
                fmt = None
                if any(h in name for h in MONEY_HINTS) and "share" not in name:
                    fmt = "#,##0"
                elif any(h in name for h in PCT_HINTS) and "revenue_per" not in name:
                    fmt = "0.0%"
                for c in col[1:]:
                    c.font = Font(name="Arial", size=10)
                    if fmt and isinstance(c.value, (int, float)):
                        c.number_format = fmt
    wb = load_workbook(path)                       # KPI_Summary mixes units: format each row by its own meaning
    for row in wb["KPI_Summary"].iter_rows(min_row=2):
        row[1].number_format = definitions[row[0].value][1]
    wb.save(path)
    print("Wrote", path)
    for p in sorted(cfg.DASHBOARD.glob("*.csv")):
        print(f"  {p.name:<48} {p.stat().st_size / 1e6:6.2f} MB")


if __name__ == "__main__":
    main()
